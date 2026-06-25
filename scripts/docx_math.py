# -*- coding: utf-8 -*-
"""正文中的 LaTeX 数学片段 → Word 原生公式(OMML)。

约定写法（行内混入中文正文）:
  $E=mc^2$  或  \\( x^2+y^2=z^2 \\)  或  \\[ \\sum_{i=1}^n x_i \\]
  显示公式可用 $$...$$（仍嵌入当前段，不另起居中段）。

依赖: pip install math2docx
"""
from __future__ import annotations

import re
from copy import deepcopy
from typing import Callable, Iterable, List, Tuple

from docx.oxml import OxmlElement
from docx.oxml.ns import qn

try:
    import math2docx
except ImportError:
    math2docx = None  # type: ignore

Part = Tuple[str, str]  # ("text"|"math", payload)

_MATH_RE = re.compile(
    r"\$\$(.+?)\$\$"
    r"|\$(?!\$)(.+?)(?<!\$)\$"
    r"|\\\[(.+?)\\\]"
    r"|\\\((.+?)\\\)",
    re.DOTALL,
)

_warned_missing = False


def split_text_math(text: str) -> List[Part]:
    """按顺序拆成 ('text', 纯文) 与 ('math', latex) 片段。"""
    if not text or not _MATH_RE.search(text):
        return [("text", text or "")]
    out: List[Part] = []
    last = 0
    for m in _MATH_RE.finditer(text):
        if m.start() > last:
            out.append(("text", text[last : m.start()]))
        latex = next(g for g in m.groups() if g is not None)
        out.append(("math", latex.strip()))
        last = m.end()
    if last < len(text):
        out.append(("text", text[last:]))
    return out


def latex_to_omml(latex: str, is_italic: bool = True):
    """LaTeX → lxml oMath 元素。"""
    if math2docx is None:
        raise ImportError("缺少 math2docx，请执行: pip install math2docx")
    return math2docx._formula(latex, is_italic=is_italic)


def _warn_missing_once():
    global _warned_missing
    if not _warned_missing:
        _warned_missing = True
        print("[docx_math 告警] 未安装 math2docx，公式将保留为 LaTeX 原文。pip install math2docx")


def append_text_run(p_el, text: str, rpr_copy=None):
    r = OxmlElement("w:r")
    if rpr_copy is not None:
        r.append(deepcopy(rpr_copy))
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    p_el.append(r)


def fill_oxml_paragraph(p_el, text: str, rpr_copy=None, is_italic: bool = True):
    """向 w:p 填入「纯文 run + oMath」混排内容（用于 build_patent 克隆段）。"""
    parts = split_text_math(text)
    has_math = any(k == "math" for k, _ in parts)
    if has_math and math2docx is None:
        _warn_missing_once()
        append_text_run(p_el, text, rpr_copy)
        return
    for kind, payload in parts:
        if kind == "text" and payload:
            append_text_run(p_el, payload, rpr_copy)
        elif kind == "math" and payload:
            try:
                p_el.append(latex_to_omml(payload, is_italic=is_italic))
            except Exception as e:
                print(f"[docx_math 告警] 公式转换失败，保留原文: {payload!r} ({e})")
                append_text_run(p_el, f"${payload}$", rpr_copy)


def fill_docx_paragraph(p, text: str, style_run: Callable | None = None, is_italic: bool = True):
    """向 python-docx Paragraph 填入混排内容（用于 build_patent_cnipa）。"""
    parts = split_text_math(text)
    has_math = any(k == "math" for k, _ in parts)
    if has_math and math2docx is None:
        _warn_missing_once()
        r = p.add_run(text)
        if style_run:
            style_run(r)
        return
    for kind, payload in parts:
        if kind == "text" and payload:
            r = p.add_run(payload)
            if style_run:
                style_run(r)
        elif kind == "math" and payload:
            try:
                math2docx.add_math(p, payload, is_italic=is_italic)
            except Exception as e:
                print(f"[docx_math 告警] 公式转换失败，保留原文: {payload!r} ({e})")
                r = p.add_run(f"${payload}$")
                if style_run:
                    style_run(r)
